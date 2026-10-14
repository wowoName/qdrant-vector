import asyncio

import pytest

from app.parsers.pdf_parser import PDFParser, PDFExtractionError


@pytest.fixture
def text_pdf(tmp_path):
    import fitz
    path = tmp_path / "integrity.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((50, 50), "Lab results\n2026\n5\nDose 5 mg\nTAIL-NOTE")
        page = doc.new_page()
        page.insert_text((50, 50), "Second page\n7")
        doc.save(path)
    return path


def test_pdf_full_stream_agree_and_short_values_survive(text_pdf):
    parser = PDFParser()
    full = parser._parse_sync(str(text_pdf))

    async def stream():
        return [page async for page in parser.parse_pages(str(text_pdf))]

    streamed = asyncio.run(stream())
    assert full.pages == [page.text for page in streamed]
    assert "2026" in full.pages[0] and "TAIL-NOTE" in full.pages[0]
    assert "7" in full.pages[1]
    from app.processors.pdf_processor import PDFProcessor
    chunks = asyncio.run(PDFProcessor().chunk(str(text_pdf)))
    assert {chunk.page_num for chunk in chunks} == {0, 1}


def test_scanned_pdf_uses_ocr_and_failure_is_visible(tmp_path, monkeypatch):
    import fitz
    from PIL import Image
    import io
    path = tmp_path / "scan.pdf"
    buf = io.BytesIO()
    Image.new("RGB", (400, 400), "white").save(buf, format="PNG")
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=buf.getvalue())
        doc.save(path)
    monkeypatch.setattr(PDFParser, "_ocr_page", lambda self, page: "OCR scan contents\n123")
    parsed = PDFParser()._parse_sync(str(path))
    assert "123" in parsed.text
    assert parsed.metadata["ocr_pages"] == [1]
    monkeypatch.setattr(PDFParser, "_ocr_page", lambda self, page: "")
    with pytest.raises(PDFExtractionError, match="OCR 未识别"):
        PDFParser()._parse_sync(str(path))


def test_empty_pdf_does_not_claim_success(tmp_path):
    import fitz
    path = tmp_path / "empty.pdf"
    with fitz.open() as doc:
        doc.new_page()
        doc.save(path)
    with pytest.raises(PDFExtractionError, match="未提取到"):
        PDFParser()._parse_sync(str(path))


def test_pdf_stream_early_close_releases_document_on_its_worker(text_pdf, monkeypatch):
    import threading
    parser = PDFParser()
    original = parser._iter_pages_sync
    closed = threading.Event()
    workers = set()

    def wrapped(path):
        iterator = original(path)
        try:
            for page in iterator:
                workers.add(threading.get_ident())
                yield page
        finally:
            workers.add(threading.get_ident())
            iterator.close()
            closed.set()

    monkeypatch.setattr(parser, "_iter_pages_sync", wrapped)

    async def run():
        stream = parser.parse_pages(str(text_pdf))
        await anext(stream)
        await stream.aclose()

    asyncio.run(run())
    assert closed.is_set() and len(workers) == 1


def test_docx_table_is_indexed_in_body_order(tmp_path):
    from docx import Document
    from app.parsers.docx_parser import DOCXParser
    doc = Document()
    doc.add_heading("First section", level=1)
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Drug"
    table.cell(0, 1).text = "Dose"
    table.cell(1, 0).text = "TEST-DRUG"
    table.cell(1, 1).text = "5 mg"
    doc.add_heading("Second section", level=1)
    doc.add_paragraph("After table")
    path = tmp_path / "table.docx"
    doc.save(path)
    parsed = DOCXParser()._parse_sync(str(path))
    assert "TEST-DRUG" in parsed.pages[0]
    assert parsed.text.index("Before table") < parsed.text.index("TEST-DRUG") < parsed.text.index("After table")


def test_csv_record_processing_and_structured_import(tmp_path):
    from app.processors.csv_processor import CSVProcessor
    from app.processors.excel_processor import ExcelProcessor
    path = tmp_path / "rows.csv"
    path.write_text("病历号,诊断\n001,肺炎\n002,\n", encoding="utf-8")
    processor = CSVProcessor()
    chunks = asyncio.run(processor.chunk(str(path)))
    records = asyncio.run(ExcelProcessor("csv").records(str(path), "doc"))
    assert len(chunks) == len(records) == 2
    assert records[1].data.get("诊断", "") == ""


def test_html_without_paragraph_tags_keeps_text(tmp_path):
    from app.parsers.html_parser import HTMLParser
    path = tmp_path / "plain.html"
    path.write_text("<html><body><div>heart <span>rate</span> 80</div></body></html>", encoding="utf-8")
    result = HTMLParser()._parse_sync(str(path))
    assert "heart" in result.text and "rate" in result.text and "80" in result.text


def test_pdf_table_keeps_field_value_relationship(tmp_path):
    import fitz
    path = tmp_path / "table.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        for x in (50, 200, 350):
            page.draw_line((x, 50), (x, 150))
        for y in (50, 100, 150):
            page.draw_line((50, y), (350, y))
        for x, y, text in [(60, 80, "Drug"), (210, 80, "Dose"), (60, 130, "TEST-DRUG"), (210, 130, "5 mg")]:
            page.insert_text((x, y), text)
        doc.save(path)
    parsed = PDFParser()._parse_sync(str(path))
    assert "Drug: TEST-DRUG" in parsed.text and "Dose: 5 mg" in parsed.text


def test_title_only_ppt_slide_and_page_slots_survive(tmp_path):
    from pptx import Presentation
    from app.parsers.ppt_parser import PPTParser
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "ONLY-TITLE"
    prs.slides.add_slide(prs.slide_layouts[6])
    third = prs.slides.add_slide(prs.slide_layouts[0])
    third.shapes.title.text = "THIRD-SLIDE"
    path = tmp_path / "slides.pptx"
    prs.save(path)
    parsed = PPTParser()._parse_sync(str(path))
    assert len(parsed.pages) == 3 and "ONLY-TITLE" in parsed.pages[0]
    assert "THIRD-SLIDE" in parsed.pages[2]
    assert parsed.pages[0].count("ONLY-TITLE") == 1


def test_markdown_code_indentation_survives(tmp_path):
    from app.parsers.markdown_parser import MarkdownParser
    path = tmp_path / "code.md"
    path.write_text("# Code\n```python\nif ok:\n    return 1\n```", encoding="utf-8")
    parsed = MarkdownParser()._parse_sync(str(path))
    assert "    return 1" in parsed.text


@pytest.mark.parametrize("mode", ["standard", "chunk_streaming", "page_streaming"])
def test_pdf_pipeline_modes_preserve_source_and_fit_model(text_pdf, monkeypatch, mode):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from app.chunkers.base import ChunkItem
    from app.processors.pdf_processor import PDFProcessor
    from app.services import document_tasks
    from app.utils.embedding_chunks import split_chunks_by_tokens

    class Tokenizer:
        def encode(self, text, **kwargs):
            return [0] + list(text) + [1]

    class Provider:
        async def prepare_chunks(self, chunks):
            return split_chunks_by_tokens(chunks, Tokenizer(), 32, overlap=3)

        async def embed_texts(self, texts):
            assert all(len(text) + 2 <= 32 for text in texts)
            return [[1.0, 0.0] for text in texts]

    class DB:
        def __init__(self):
            self.saved = []

        def add(self, item):
            item.id = f"chunk-{len(self.saved)}"
            self.saved.append(item)

        async def flush(self):
            pass

        async def commit(self):
            pass

        async def execute(self, statement):
            return SimpleNamespace(scalar_one=lambda: SimpleNamespace(directory_id=None, name="test.pdf"))

    provider = Provider()
    monkeypatch.setattr(document_tasks, "get_embedding", AsyncMock(return_value=provider))
    db = DB()
    qdrant = SimpleNamespace(upsert_points=AsyncMock())
    service = SimpleNamespace(check_cancelled=AsyncMock(), update_progress=AsyncMock())
    processor = PDFProcessor()
    version = SimpleNamespace(id="version", created_at=None)
    common = dict(db=db, qdrant=qdrant, task_id="task", task_svc=service,
                  version=version, document_id="doc", collection_id="collection",
                  file_type="pdf", file_path=str(text_pdf), processor=processor)

    async def run():
        if mode == "page_streaming":
            return await document_tasks._process_page_streaming(**common, total_pages=2, file_size=text_pdf.stat().st_size)
        parsed = await processor.parse(str(text_pdf))
        if mode == "standard":
            return await document_tasks._process_standard(**common, parse_result=parsed)
        return await document_tasks._process_chunk_streaming(**common, parse_result=parsed, text_len=len(parsed.text))

    count = asyncio.run(run())
    assert count == len(db.saved) and count > 2
    assert [chunk.chunk_index for chunk in db.saved] == list(range(count))
    assert {chunk.page_num for chunk in db.saved} == {0, 1}
    assert any("TAIL-NOTE" in chunk.text for chunk in db.saved)
    assert any("2026" in chunk.text for chunk in db.saved)
