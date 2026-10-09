"""
Excel 解析器
"""

import asyncio
from typing import List

from .base import BaseParser, ParseResult
from app.core.logging import get_logger
from app.utils.text_cleaner import normalize_text

logger = get_logger(__name__)


class ExcelParser(BaseParser):
    """Excel 文档解析器"""

    def supported_types(self) -> List[str]:
        return ["xlsx", "xls", "excel"]

    async def parse(self, file_path: str) -> ParseResult:
        logger.info(f"Parsing Excel: {file_path}")
        return await asyncio.to_thread(self._parse_sync, file_path)

    def _parse_sync(self, file_path: str) -> ParseResult:
        import os
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".xls":
            return self._parse_xls(file_path)
        return self._parse_xlsx(file_path)

    def _parse_xlsx(self, file_path: str) -> ParseResult:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True, read_only=True)
        pages = []  # 每个 sheet 一页
        all_tables = []
        all_text_parts = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_parts = [f"# Sheet: {sheet_name}"]
            table_data = []
            row_num = 0

            for row in ws.iter_rows(values_only=True):
                # 跳过全空行
                if all(v is None or str(v).strip() == "" for v in row):
                    continue

                row_strs = ["" if v is None else str(v).strip() for v in row]
                table_data.append(row_strs)

                # 转换为描述性文本
                if row_num == 0:
                    # 表头
                    headers = row_strs
                    sheet_parts.append("表头: " + " | ".join(headers))
                else:
                    # 数据行：列名: 值
                    line_parts = []
                    for i, val in enumerate(row_strs):
                        if val and i < len(headers):
                            line_parts.append(f"{headers[i]}: {val}")
                    if line_parts:
                        sheet_parts.append("; ".join(line_parts))

                row_num += 1

            if len(sheet_parts) > 1:
                sheet_text = "\n".join(sheet_parts)
                pages.append(sheet_text)
                all_text_parts.append(sheet_text)
            all_tables.append(table_data)

        wb.close()

        full_text = "\n\n".join(all_text_parts)
        full_text = normalize_text(full_text)

        return ParseResult(
            text=full_text,
            pages=pages,
            metadata={"sheet_count": len(wb.sheetnames), "sheets": wb.sheetnames},
            tables=all_tables,
        )

    def _parse_xls(self, file_path: str) -> ParseResult:
        """解析旧版 .xls 文件（Excel 97-2003）"""
        import os
        try:
            import xlrd
        except ImportError:
            logger.warning("xlrd not installed, cannot parse .xls files")
            return ParseResult(
                text="",
                pages=[],
                metadata={"error": "xlrd not installed. Install xlrd==1.2.0 for .xls support."},
            )

        try:
            wb = xlrd.open_workbook(file_path)
        except Exception as e:
            logger.error(f"Failed to open .xls file: {e}")
            # 回退：尝试 LibreOffice 转换
            return self._convert_with_libreoffice(file_path)

        pages = []
        all_tables = []
        all_text_parts = []

        for sheet_idx in range(wb.nsheets):
            ws = wb.sheet_by_index(sheet_idx)
            sheet_name = wb.sheet_names()[sheet_idx]
            sheet_parts = [f"# Sheet: {sheet_name}"]
            table_data = []

            if ws.nrows == 0:
                continue

            headers = []
            for row_num in range(ws.nrows):
                row_values = []
                for col_num in range(ws.ncols):
                    cell = ws.cell(row_num, col_num)
                    val = ""
                    if cell.ctype == xlrd.XL_CELL_NUMBER:
                        # 整数就显示为整数
                        if cell.value == int(cell.value):
                            val = str(int(cell.value))
                        else:
                            val = str(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            from datetime import datetime
                            dt = xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode)
                            val = dt.strftime("%Y-%m-%d %H:%M:%S")
                        except Exception:
                            val = str(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        val = str(bool(cell.value))
                    else:
                        val = str(cell.value).strip() if cell.value else ""
                    row_values.append(val)

                # 跳过全空行
                if all(v == "" or v is None for v in row_values):
                    continue

                table_data.append(row_values)

                if row_num == 0:
                    headers = row_values
                    sheet_parts.append("表头: " + " | ".join(headers))
                else:
                    line_parts = []
                    for i, val in enumerate(row_values):
                        if val and i < len(headers):
                            line_parts.append(f"{headers[i]}: {val}")
                    if line_parts:
                        sheet_parts.append("; ".join(line_parts))

            if len(sheet_parts) > 1:
                sheet_text = "\n".join(sheet_parts)
                pages.append(sheet_text)
                all_text_parts.append(sheet_text)
            all_tables.append(table_data)

        full_text = "\n\n".join(all_text_parts)
        full_text = normalize_text(full_text)

        return ParseResult(
            text=full_text,
            pages=pages,
            metadata={"sheet_count": wb.nsheets, "sheets": wb.sheet_names(), "format": "xls"},
            tables=all_tables,
        )

    def _convert_with_libreoffice(self, file_path: str) -> ParseResult:
        """用 LibreOffice 把 .xls 转成 .xlsx 再解析"""
        import os
        import tempfile
        import subprocess
        import shutil

        tmp_dir = tempfile.mkdtemp(prefix="xls_convert_")
        try:
            result = subprocess.run(
                ["soffice", "--headless", "--convert-to", "xlsx",
                 "--outdir", tmp_dir, file_path],
                capture_output=True,
                text=True,
                timeout=120,
            )
            if result.returncode != 0:
                return ParseResult(
                    text="", pages=[],
                    metadata={"error": f"LibreOffice conversion failed: {result.stderr[:200]}"},
                )

            base_name = os.path.splitext(os.path.basename(file_path))[0]
            converted = os.path.join(tmp_dir, f"{base_name}.xlsx")
            if not os.path.exists(converted):
                for f in os.listdir(tmp_dir):
                    if f.endswith(".xlsx"):
                        converted = os.path.join(tmp_dir, f)
                        break

            if os.path.exists(converted):
                parsed = self._parse_xlsx(converted)
                if parsed.metadata:
                    parsed.metadata["converter"] = "libreoffice"
                return parsed
        except FileNotFoundError:
            logger.debug("LibreOffice not available")
        except Exception as e:
            logger.debug(f"LibreOffice conversion failed: {e}")
        finally:
            try:
                shutil.rmtree(tmp_dir, ignore_errors=True)
            except Exception:
                pass

        return ParseResult(
            text="", pages=[],
            metadata={"error": "No .xls parser available. Install xlrd or LibreOffice."},
        )
