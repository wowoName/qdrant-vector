"""
查询明细导出服务

当查询结果的明细超过内联展示上限（30 条）时，
将完整明细导出为 Excel 文件，供用户通过下载链接获取。
"""

import os
import re
import time
from pathlib import Path
from uuid import uuid4

from app.core.config import settings
from app.core.logging import get_logger
from app.utils.datetime_utils import now

logger = get_logger(__name__)

# 导出文件保留时长（24 小时，过期自动清理）
EXPORT_TTL_SECONDS = 24 * 3600

# 文件名安全校验（防路径穿越）
_FILENAME_RE = re.compile(r"^[\w\u4e00-\u9fa5\-]+\.xlsx$")


def get_export_dir() -> Path:
    """导出目录：data/exports（与 storage 同级）"""
    export_dir = Path(settings.storage_root).resolve().parent / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    return export_dir


def cleanup_expired_exports() -> None:
    """清理超过保留时长的导出文件（尽力而为，不影响主流程）"""
    try:
        export_dir = get_export_dir()
        now = time.time()
        for f in export_dir.glob("*.xlsx"):
            try:
                if now - f.stat().st_mtime > EXPORT_TTL_SECONDS:
                    f.unlink()
            except OSError:
                pass
    except Exception as e:
        logger.warning(f"清理过期导出文件失败: {e}")


def export_sheets_to_excel(
    sheets: list[tuple[str, list[str], list[list]]],
) -> str | None:
    """导出多个 sheet 到 Excel 文件

    Args:
        sheets: [(sheet名, 列名列表, 数据行列表), ...]

    Returns:
        导出文件名；失败返回 None
    """
    if not sheets:
        return None

    try:
        from openpyxl import Workbook

        try:
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        except ImportError:
            ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")

        def _clean(v) -> str:
            """清理 Excel 非法字符"""
            s = "" if v is None else str(v)
            return ILLEGAL_CHARACTERS_RE.sub("", s)

        wb = Workbook()
        first = True
        for sheet_name, columns, rows in sheets:
            if not columns or not rows:
                continue
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = (sheet_name or "明细数据")[:31]
            ws.append([_clean(c) for c in columns])
            for row in rows:
                ws.append([_clean(row[i]) if i < len(row) else "" for i in range(len(columns))])

            # 表头加粗
            from openpyxl.styles import Font
            for cell in ws[1]:
                cell.font = Font(bold=True)

            # 合理的列宽（按列名和前几行内容估算，上限 40）
            for i, col in enumerate(columns):
                max_len = len(str(col))
                for row in rows[:50]:
                    if i < len(row):
                        max_len = max(max_len, len(str(row[i] or "")))
                col_letter = ws.cell(row=1, column=i + 1).column_letter
                ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        if first:
            # 没有有效 sheet
            return None

        # 文件名：明细导出_时间戳_随机.xlsx
        filename = f"明细导出_{now():%Y%m%d_%H%M%S}_{uuid4().hex[:6]}.xlsx"
        file_path = get_export_dir() / filename
        wb.save(str(file_path))

        # 顺手清理过期文件
        cleanup_expired_exports()

        total_rows = sum(len(r) for _, _, r in sheets)
        logger.info(f"明细已导出: {filename}（{len(sheets)} 个 sheet，共 {total_rows} 行）")
        return filename

    except Exception as e:
        logger.warning(f"明细导出失败: {e}")
        return None


def export_rows_to_excel(
    columns: list[str],
    rows: list[list],
    title: str = "明细数据",
) -> str | None:
    """将行列数据导出为 Excel 文件（单 sheet）

    Args:
        columns: 列名列表
        rows: 数据行列表（每行为与 columns 对应的值列表）
        title: 文件标题（用于 sheet 名）

    Returns:
        导出文件名（用于拼接下载链接）；失败返回 None
    """
    if not columns or not rows:
        return None
    return export_sheets_to_excel([(title, columns, rows)])


def resolve_export_file(filename: str) -> Path | None:
    """根据文件名解析导出文件路径（含安全校验）"""
    if not filename or not _FILENAME_RE.match(filename):
        return None
    file_path = get_export_dir() / filename
    if file_path.is_file():
        return file_path
    return None
